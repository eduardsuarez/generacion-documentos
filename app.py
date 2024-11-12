from uuid import uuid4
import uuid
import openpyxl
import docx
import os
from datetime import datetime
from fastapi import FastAPI, Request
from pydantic import BaseModel

app = FastAPI()

@app.get("/")
async def home():
    return {"Message": "Bienvenido a mi API"}

#Remmplazar valores en plantilla .docx
@app.post("/replaceDocx/")
async def reemplazar_docx(ruta_documento, datos):
    # Reemplazar valores en la plantilla
    ruta_actual = os.getcwd()
    ruta_plantilla = os.path.join(ruta_actual,"plantillas","formato-documento.docx")
    datos = {
        "FECHA_ACTUAL" : datetime.now().strftime("%H:%M:%S"),
        "INDUSTRIA_ESPECIFICA": "Redes de telecomunicaciones",
        "METRICA_RENDIMIENTO": "ROI",
        "NOMBRE_PLATAFORMA": "TELCO-BIT"
    }
    documento = docx.Document(ruta_documento)

    for paragraph in documento.paragraphs:
        for run in paragraph.runs:
            for key, value in datos.items():
                if f"${{{key}}}" in run.text:
                    run.text = run.text.replace(f"${{{key}}}", value)
    ruta_salida = f"temp_{uuid.uuid4()}.docx"
    ruta_salida_tmp = os.path.join(ruta_actual,"tmp",ruta_salida)
    documento.save(ruta_salida_tmp)


# Reemplazar valores en la plantilla excel
app.post("/replaceXlsx/")
async def reemplazar_xlsx(ruta_plantilla, valores):
    wb = openpyxl.load_workbook(ruta_plantilla)
    ws = wb.active
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "${" in cell.value:
                # Reemplazar el marcador con el valor correspondiente
                for clave, valor in valores.items():
                    if f"${{{clave}}}" in cell.value:
                        cell.value = cell.value.replace(f"${{{clave}}}", str(valor))
    ruta_salida = f"{uuid.uuid4}.xlsx"
    wb.save(ruta_salida)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=4449)
